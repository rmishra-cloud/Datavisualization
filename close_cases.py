import argparse
import os
import sys
import time
from typing import List

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.chrome.options import Options as ChromeOptions
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager


DEFAULT_URL = "https://www.coursefinder.ai/Application"
DEFAULT_CLOSE_COMMENT = (
    "Dear Team,\n"
    "We are closing the captioned student case as per the discussion on the offer follow-up sheet. \n\n"
    "Thanks and Regards\n"
)


def read_ack_list(excel_path: str, preferred_column: str) -> List[str]:
    if not os.path.isfile(excel_path):
        raise FileNotFoundError(f"Excel not found: {excel_path}")

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook.active

    # Determine headers from the first non-empty row
    headers = []
    header_row_index = None
    for row in sheet.iter_rows(min_row=1, max_row=10, values_only=True):
        if row and any(cell is not None and str(cell).strip() for cell in row):
            headers = ["" if cell is None else str(cell).strip() for cell in row]
            header_row_index = headers and 1 or 1
            break

    if not headers:
        raise ValueError("Could not detect header row in Excel file.")

    candidate_columns = [
        preferred_column,
        "Ack No",
        "AckNo",
        "Acknowledgement",
        "AcknowledgementNumber",
        "Acknowledgement No",
        "Acknowledgement_Number",
        "Acknowledgement_No",
        "Acknowledgement ID",
        "AcknowledgementID",
    ]

    found_index = None
    for candidate in candidate_columns:
        try:
            found_index = headers.index(candidate)
            break
        except ValueError:
            continue

    if found_index is None:
        raise KeyError(
            f"Could not find an ACK column. Looked for: {candidate_columns}. Available headers: {headers}"
        )

    ack_list: List[str] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or len(row) <= found_index:
            continue
        value = row[found_index]
        if value is None:
            continue
        value_str = str(value).strip()
        if value_str:
            ack_list.append(value_str)

    if not ack_list:
        raise ValueError("ACK list is empty after reading the Excel file.")

    return ack_list


def build_webdriver(headless: bool) -> webdriver.Chrome:
    options = ChromeOptions()
    if headless:
        # Using new headless for modern Chrome
        options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--window-size=1920,1080")

    driver = webdriver.Chrome(
        service=Service(ChromeDriverManager().install()),
        options=options,
    )
    return driver


def process_acknowledgements(
    excel_path: str,
    url: str,
    ack_column: str,
    close_comment: str,
    headless: bool,
    wait_seconds: int,
) -> None:
    ack_list = read_ack_list(excel_path, ack_column)
    print(f"Loaded {len(ack_list)} ACK(s) from: {excel_path}")

    driver = build_webdriver(headless=headless)
    driver.maximize_window()
    driver.get(url)

    wait = WebDriverWait(driver, wait_seconds)

    try:
        # Wait for the ACK input to be present on the page
        wait.until(EC.presence_of_element_located((By.XPATH, '//*[@id="AcknowledgementNumber"]')))
    except Exception as error:
        driver.quit()
        raise RuntimeError(
            "Could not find the ACK input field. Ensure you are logged in and the page has loaded."
        ) from error

    skipped_statuses = {
        "Payment Received",
        "Application on Hold (Reason Required)",
        "Deferral (Reason Required)",
        "Intake yet to open",
    }

    for ack in ack_list:
        try:
            ack_input = wait.until(
                EC.element_to_be_clickable((By.XPATH, '//*[@id="AcknowledgementNumber"]'))
            )
            ack_input.clear()
            ack_input.send_keys(str(ack))
            ack_input.send_keys(Keys.TAB)

            search_button = wait.until(
                EC.element_to_be_clickable((By.XPATH, '//*[@id="filterApplications"]'))
            )
            search_button.click()

            # Wait for status dropdown to be ready
            status_dropdown_element = wait.until(
                EC.presence_of_element_located((By.XPATH, '//*[@id="Declaration_StageId"]'))
            )
            status_dropdown = Select(status_dropdown_element)
            current_status = status_dropdown.first_selected_option.text.strip()

            if current_status in skipped_statuses:
                print(f"Skipped (status = {current_status}) for ACK: {ack}")
                continue

            # Select Case Closed
            status_dropdown.select_by_visible_text("Case Closed (Reason Required)")

            # Reason dropdown appears/updates after selecting status
            reason_dropdown_element = wait.until(
                EC.presence_of_element_located((By.XPATH, '//*[@id="Declaration_StageId_child"]'))
            )
            reason_dropdown = Select(reason_dropdown_element)
            reason_dropdown.select_by_visible_text(
                "Offer Received - Student not interested to pay"
            )

            comment_box = wait.until(
                EC.element_to_be_clickable((By.XPATH, '//*[@id="txtMsg"]'))
            )
            comment_box.clear()
            comment_box.send_keys(close_comment)

            # Submit and send
            submit_button = wait.until(
                EC.element_to_be_clickable((By.XPATH, '//*[@id="btn-special-user"]'))
            )
            submit_button.click()

            send_button = wait.until(
                EC.element_to_be_clickable((By.XPATH, '//*[@id="btnSend"]'))
            )
            send_button.click()

            print(f"Case closed for ACK: {ack}")

        except Exception as error:
            print(f"Error for ACK {ack}: {error}")

    print("All ACKs processed.")
    driver.quit()


def parse_arguments(argv: List[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Close cases in coursefinder.ai by ACK.")
    parser.add_argument(
        "--excel-path",
        default=os.environ.get("EXCEL_PATH", ""),
        help="Path to the Excel file containing an ACK column.",
    )
    parser.add_argument(
        "--ack-column",
        default=os.environ.get("ACK_COLUMN", "Ack No"),
        help="Column name in Excel that contains ACK numbers.",
    )
    parser.add_argument(
        "--url",
        default=os.environ.get("COURSEFINDER_URL", DEFAULT_URL),
        help="Application URL.",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run Chrome in headless mode.",
    )
    parser.add_argument(
        "--wait-seconds",
        type=int,
        default=int(os.environ.get("SELENIUM_WAIT_SECONDS", "20")),
        help="Explicit wait timeout in seconds.",
    )
    parser.add_argument(
        "--comment",
        default=os.environ.get("CLOSE_COMMENT", DEFAULT_CLOSE_COMMENT),
        help="Close comment text.",
    )
    return parser.parse_args(argv)


def main(argv: List[str]) -> int:
    args = parse_arguments(argv)

    if not args.excel_path:
        print("Error: --excel-path is required (or set EXCEL_PATH env var).", file=sys.stderr)
        return 2

    try:
        process_acknowledgements(
            excel_path=args.excel_path,
            url=args.url,
            ack_column=args.ack_column,
            close_comment=args.comment,
            headless=args.headless,
            wait_seconds=args.wait_seconds,
        )
        return 0
    except Exception as error:
        print(f"Failed: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))

